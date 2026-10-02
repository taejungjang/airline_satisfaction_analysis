import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from config import *
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

D = str(ROOT) + '/'
NAVY = RGBColor(0x1F, 0x3A, 0x5F); ORG = RGBColor(0xE0, 0x7B, 0x39); GRY = RGBColor(0x5B, 0x66, 0x73)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8); WHITE = RGBColor(255, 255, 255); INK = RGBColor(0x22, 0x2B, 0x36)
MID = RGBColor(0xC9, 0xD1, 0xDB)
FONT = 'Malgun Gothic'
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def tb(s, x, y, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = b.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(4)
        r = p.add_run(); r.text = ln; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return b


def rect(s, x, y, w, h, fill):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill; r.line.fill.background(); r.shadow.inherit = False
    return r


STAGES = ['현황', '문제 정의', '분석', '제안']


def slide(title, n, stage):
    s = prs.slides.add_slide(blank)
    for i, st in enumerate(STAGES):
        on = i == stage
        x = 7.85 + i * 1.3
        rect(s, x, 0.3, 1.2, 0.34, NAVY if on else LIGHT)
        tb(s, x, 0.3, 1.2, 0.34, st, 11, on, WHITE if on else GRY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, 0.6, 0.85, 12.1, 0.9, title, 26, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 0.65, 1.75, 0.9, 0.06, ORG)
    tb(s, 12.2, 7.0, 0.9, 0.3, str(n), 11, False, GRY, PP_ALIGN.RIGHT)
    return s


def big(s, x, y, w, num, label, color=ORG, h=1.7):
    rect(s, x, y, w, h, LIGHT)
    tb(s, x, y + 0.12, w, 0.95, num, 40, True, color, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, x + 0.1, y + 1.05, w - 0.2, h - 1.05, label, 14, False, GRY, PP_ALIGN.CENTER)


def note(s, text, y=6.35):
    rect(s, 0.65, y, 12.0, 0.65, NAVY)
    tb(s, 0.85, y, 11.6, 0.65, text, 17, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)


def defbox(s, text, y=1.95, h=0.6):
    rect(s, 0.65, y, 12.0, h, LIGHT); rect(s, 0.65, y, 0.07, h, ORG)
    tb(s, 0.9, y, 11.7, h, text, 14, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)


DEF = '기준: 서비스 평가가 4점 미만 → 4점 이상으로 오를 때, 전체 비행 만족률이 몇 %p 오르는가'

# 1 표지
s = prs.slides.add_slide(blank); rect(s, 0, 0, 13.333, 7.5, NAVY); rect(s, 0.8, 3.55, 1.2, 0.07, ORG)
tb(s, 0.8, 1.7, 11.5, 1.6, '항공 승객 만족도 분석', 44, True, WHITE)
tb(s, 0.8, 3.8, 11.5, 1.2, ['한정된 예산, 무엇부터 개선해야 만족도가 오를까?', '설문 103,904건 기반 개선 우선순위 도출'], 20, False, RGBColor(0xDD, 0xE4, 0xEE))
tb(s, 0.8, 6.3, 11.8, 0.5, 'Taejung Jang  |  Python · statsmodels · scikit-learn  |  Kaggle Airline Passenger Satisfaction', 13, False, RGBColor(0xAA, 0xB8, 0xC8))

# 2 현황
s = slide('승객 10명 중 6명은 전체 비행에 만족하지 못함', 2, 0)
tb(s, 0.65, 1.95, 12, 0.4, '※ Kaggle 항공 승객 만족도 설문 데이터 기준 (승객 103,904명)', 15, True, ORG)
tb(s, 0.65, 2.5, 12, 0.4, '전체 비행에 대한 종합 만족도', 15, False, GRY)
W = 12.0
rect(s, 0.65, 3.0, W * 0.433, 1.3, ORG); rect(s, 0.65 + W * 0.433, 3.0, W * 0.567, 1.3, MID)
tb(s, 0.65, 3.0, W * 0.433, 1.3, '만족 43%', 30, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
tb(s, 0.65 + W * 0.433, 3.0, W * 0.567, 1.3, '중립·불만 57%', 30, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
tb(s, 0.65, 4.9, 12, 1.5, ['전체 비행 만족률을 높이려면 어떤 서비스부터 개선해야 하는가?'], 24, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# 3 문제 정의
s = slide('핵심 질문: 14개 서비스 중 어디에 먼저 투자할 것인가?', 3, 1)
rect(s, 0.65, 2.0, 12.0, 1.1, LIGHT)
tb(s, 0.65, 2.0, 12.0, 1.1, '예산은 한정됨 → 서비스별 개선 효과를 숫자로 비교해야 함', 22, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
rect(s, 0.65, 3.3, 12.0, 1.7, LIGHT); rect(s, 0.65, 3.3, 0.1, 1.7, ORG)
tb(s, 1.0, 3.35, 11.5, 0.45, '정량 판단 기준', 14, True, ORG, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 1.0, 3.8, 11.5, 0.65, '서비스를 개선했을 때, 전체 비행 만족률이 몇 %p 오르는가', 22, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 1.0, 4.45, 11.5, 0.5, '예) 서비스 A 개선 +○%p  vs  서비스 B 개선 +○%p  →  큰 쪽에 우선 투자', 14, False, GRY, anchor=MSO_ANCHOR.MIDDLE)
for i, (h, b_) in enumerate([('1. 정량화', '서비스별 만족률 상승폭(%p) 산출'), ('2. 우선순위', '상승폭 큰 순서로 정렬'), ('3. 추가 확인', '승객 유형(출장/개인 여행)별 차이')]):
    x = 0.65 + i * 4.1
    rect(s, x, 5.25, 3.8, 1.5, NAVY if i < 2 else GRY)
    tb(s, x, 5.3, 3.8, 0.6, h, 19, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, x + 0.1, 5.95, 3.6, 0.7, b_, 14, False, WHITE, PP_ALIGN.CENTER)
tb(s, 0.65, 6.85, 12, 0.35, '※ "개선"의 구체적 기준은 데이터 분석 후 정의', 11, False, GRY)

# 4 데이터
s = slide('분석 데이터: Kaggle 항공 승객 만족도 설문', 4, 2)
tb(s, 0.65, 1.95, 12, 0.4, '출처: Kaggle "Airline Passenger Satisfaction"  |  승객 103,904명 × 24개 컬럼', 15, True, ORG)
rect(s, 0.65, 2.5, 5.9, 4.1, LIGHT)
tb(s, 0.85, 2.58, 5.5, 0.5, '서비스 평가 컬럼 (14개): 점수 0~5', 17, True, NAVY)
labs = [('0', '해당없음', MID, NAVY), ('1', '매우 불만', RGBColor(0xB7, 0xC3, 0xD1), NAVY), ('2', '', RGBColor(0x8F, 0xA3, 0xBA), WHITE), ('3', '보통', RGBColor(0x5F, 0x7B, 0x9C), WHITE), ('4', '', RGBColor(0x2E, 0x4C, 0x75), WHITE), ('5', '매우 만족', ORG, WHITE)]
for i, (n, t, c, tc) in enumerate(labs):
    x = 0.85 + i * 0.88
    rect(s, x, 3.2, 0.82, 0.85, c); tb(s, x, 3.2, 0.82, 0.85, n, 24, True, tc, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, x - 0.1, 4.1, 1.02, 0.5, t, 11, False, GRY, PP_ALIGN.CENTER)
tb(s, 0.85, 4.7, 5.5, 1.9, ['• 승객이 항목별로 0~5점 직접 평가', '• 점수가 높을수록 만족', '• 0점은 "해당없음" (예: 와이파이 미이용)'], 15, False, INK)
rect(s, 6.8, 2.5, 5.85, 1.95, LIGHT)
tb(s, 7.0, 2.58, 5.5, 0.5, '서비스 평가 컬럼 목록', 17, True, NAVY)
tb(s, 7.0, 3.05, 5.5, 1.2, ['출발 전: 온라인 예약 · 체크인 · 온라인 탑승 · 게이트 위치 · 출·도착 시간', '기내: 와이파이 · 좌석 · 다리 공간 · 엔터테인먼트 · 음식 · 청결 · 탑승 서비스 · 기내 서비스', '기타: 수하물 처리'], 12, False, INK)
rect(s, 6.8, 4.55, 5.85, 1.05, LIGHT)
tb(s, 7.0, 4.58, 5.5, 0.4, '승객 정보 컬럼', 17, True, NAVY)
tb(s, 7.0, 4.98, 5.5, 0.6, '성별 · 나이 · 고객 유형(충성/비충성) · 여행 목적 · 좌석 등급 · 비행 거리 · 지연 시간', 12, False, INK)
rect(s, 6.8, 5.7, 5.85, 0.9, LIGHT)
tb(s, 7.0, 5.72, 5.5, 0.4, '결과 컬럼 (분석 대상)', 17, True, ORG)
tb(s, 7.0, 6.1, 5.5, 0.45, '전체 비행 만족도: 만족 / 중립·불만', 13, False, INK)

# 5 유형별 격차
s = slide('승객 유형별 비행 만족률 격차 큼', 5, 2)
s.shapes.add_picture(D + 'charts/segments.png', Inches(0.9), Inches(1.95), width=Inches(11.5))
note(s, '→ 승객 유형(여행 목적·고객 유형·좌석 등급)을 함께 고려한 분석 필요', 6.3)

# 6 전처리
s = slide('데이터 전처리', 6, 2)
cols = [('항목', 0.65, 2.6), ('문제', 3.35, 4.4), ('처리', 7.85, 4.8)]
for h, x, w in cols:
    rect(s, x, 2.0, w - 0.05, 0.6, NAVY); tb(s, x, 2.0, w - 0.05, 0.6, h, 16, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
rows = [('결측치', '도착 지연 시간 310건(0.3%) 비어 있음', '출발 지연 시간과 상관 0.97\n→ 회귀식으로 값 추정해 채움'),
        ('이상치', '비행 거리에 극단적으로 큰 값 존재', '상·하한(IQR 기준)으로 보정\n→ 삭제하지 않아 표본 유지'),
        ('중복 정보 변수', '출발 지연과 도착 지연이 거의 같은 정보\n(다중공선성, VIF 14.9)', '도착 지연 변수 제외\n→ 서비스 효과 추정 왜곡 방지')]
for i, r in enumerate(rows):
    y = 2.75 + i * 1.3
    for (h, x, w), v in zip(cols, r):
        rect(s, x, y, w - 0.05, 1.2, LIGHT)
        tb(s, x + 0.1, y, w - 0.25, 1.2, v.split('\n'), 18 if x == 0.65 else 15, x == 0.65, NAVY if x == 0.65 else INK, PP_ALIGN.CENTER if x == 0.65 else PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

# 7 점수와 전체 만족의 관계
s = slide('서비스 평가 점수가 4점부터 전체 비행 만족률이 급증', 7, 2)
defbox(s, '서비스 평가 점수(개별 서비스)와 전체 비행 만족 여부(종합)는 별개의 설문 항목')
s.shapes.add_picture(D + 'charts/wifi.png', Inches(0.65), Inches(2.75), width=Inches(6.1))
big(s, 7.2, 2.75, 5.45, '25% → 60%', '와이파이를 3점 준 승객 vs 4점 준 승객 중\n전체 비행에 만족한 비율', ORG, 1.9)
rect(s, 7.2, 4.8, 5.45, 1.75, NAVY)
tb(s, 7.4, 4.8, 5.1, 1.75, ['1~3점은 25~33%로 비슷, 4점부터 급증', '→ 개선 기준: 4점 미만 → 4점 이상', '→ 판단 기준: 이때 전체 비행 만족률 상승폭 (%p)'], 15, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 0.65, 6.65, 7, 0.4, '※ 0점(해당없음) 제외. 다른 서비스도 비슷한 패턴', 12, False, GRY)

# 8 분석 방법
s = slide('분석 방법: 로지스틱 회귀', 8, 2)
tb(s, 0.65, 1.9, 12, 0.45, '서비스 평가 점수(입력) → 전체 비행 만족 확률(결과)의 관계를 수식으로 추정', 16, True, GRY)
rect(s, 0.65, 2.45, 12.0, 1.75, LIGHT); rect(s, 0.65, 2.45, 0.1, 1.75, ORG)
tb(s, 1.0, 2.5, 11.5, 0.8, 'ln( p / (1 − p) ) = a₁·와이파이 + a₂·온라인 탑승 + … + b·승객 유형 + c', 24, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 1.0, 3.3, 11.5, 0.85, ['p: 전체 비행 만족 확률   |   서비스 변수: 4점 이상이면 1, 아니면 0   |   승객 유형: 여행 목적·좌석 등급 등', 'a: 서비스별 효과 크기 (클수록 만족 확률을 더 크게 올림)'], 13, False, INK, anchor=MSO_ANCHOR.MIDDLE)
reasons = [('1. 관계 추정', ['서비스 점수와 만족 여부의', '관계를 수식으로 추정']),
           ('2. 다른 조건 고정', ['승객 유형이 같다고 두고', '서비스별 순수 효과 비교']),
           ('3. 쉬운 해석', ['"4점 이상 개선 시 만족률', '몇 %p 상승"으로 바로 설명'])]
for i, (h, b_) in enumerate(reasons):
    x = 0.65 + i * 4.1
    rect(s, x, 4.4, 3.8, 1.6, LIGHT); rect(s, x, 4.4, 3.8, 0.08, ORG)
    tb(s, x, 4.55, 3.8, 0.5, h, 18, True, NAVY, PP_ALIGN.CENTER)
    tb(s, x + 0.1, 5.1, 3.6, 0.9, b_, 14, False, INK, PP_ALIGN.CENTER)
note(s, '→ 이 방법으로 서비스별 "4점 이상 개선 시 전체 비행 만족률 상승폭"을 산출', 6.25)

# 9 모델 검증
s = slide('모델 검증: 만족 여부를 89% 맞혀 분석 결과 신뢰 가능', 9, 2)
big(s, 0.65, 2.0, 5.9, '89.3%', '정확도: 검증 데이터 승객의 전체 비행 만족 여부를 맞힌 비율', ORG, 2.1)
big(s, 6.75, 2.0, 5.9, '0.943', 'AUC: 만족 / 불만 승객을 구분하는 능력 (1에 가까울수록 우수)', ORG, 2.1)
rect(s, 0.65, 4.35, 12.0, 1.5, LIGHT); rect(s, 0.65, 4.35, 0.07, 1.5, NAVY)
tb(s, 0.9, 4.35, 11.6, 1.5, ['서비스 점수를 "4점 이상 여부"로 바꿔 입력했을 때 성능 개선', '정확도 87.7% → 89.3%  ·  AUC 0.928 → 0.943'], 18, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
note(s, '→ 모델이 실제 만족 여부를 잘 설명하므로, 이어지는 분석 결과를 신뢰할 수 있음', 6.1)
tb(s, 0.65, 6.85, 12, 0.35, '※ 학습 80% / 검증 20% 분할, 검증 데이터 기준', 11, False, GRY)

# 10 분석 결과 1
s = slide('분석 결과: 와이파이·온라인 탑승 개선 효과가 가장 큼', 10, 2)
defbox(s, DEF)
s.shapes.add_picture(D + 'charts/lift.png', Inches(0.6), Inches(2.7), width=Inches(6.3))
rect(s, 8.3, 2.8, 4.35, 1.5, LIGHT); rect(s, 8.3, 2.8, 0.07, 1.5, ORG)
tb(s, 8.5, 2.8, 4.0, 1.5, ['1위 와이파이 +11.8%p', '2위 온라인 탑승 +8.2%p'], 19, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 8.3, 4.5, 4.35, 1.2, LIGHT); rect(s, 8.3, 4.5, 0.07, 1.2, GRY)
tb(s, 8.5, 4.5, 4.0, 1.2, ['나머지는 +4%p 이하', '음식·게이트 위치는 영향 거의 없음'], 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)
note(s, '→ 전체 승객 기준 우선 투자 대상: 와이파이, 온라인 탑승', 6.5)

# 11 추가 확인
s = slide('추가 확인: 와이파이는 모든 승객에 효과, 온라인 탑승은 출장 중심', 11, 2)
defbox(s, '같은 기준을 출장 승객 / 개인 여행 승객으로 나눠서 확인 (고객군별 모델)')
s.shapes.add_picture(D + 'charts/seg_effect.png', Inches(0.6), Inches(2.65), width=Inches(6.6))
rect(s, 8.1, 2.8, 4.55, 1.55, LIGHT); rect(s, 8.1, 2.8, 0.07, 1.55, ORG)
tb(s, 8.3, 2.8, 4.25, 1.55, ['공통', '와이파이는 두 유형 모두 상승 (개인 여행에서 특히 큼)'], 15, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 8.1, 4.5, 4.55, 1.55, LIGHT); rect(s, 8.1, 4.5, 0.07, 1.55, NAVY)
tb(s, 8.3, 4.5, 4.25, 1.55, ['차이', '온라인 탑승·체크인은 출장에서만 상승, 개인 여행은 변화 없음'], 15, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 0.65, 6.6, 12, 0.4, '※ 서비스를 승객 유형별로 따로 제공하자는 의미가 아니라, 개선 효과가 어느 승객에게서 나타나는지 확인한 것', 11, False, GRY)

# 12 결론
s = slide('결론: 와이파이 → 온라인 탑승 순으로 우선 투자', 12, 3)
for h, x, wd in zip(['순위', '서비스', '전체 비행 만족률 상승폭', '참고'], [0.65, 1.95, 5.15, 8.15], [1.25, 3.1, 2.95, 4.5]):
    rect(s, x, 2.0, wd, 0.6, NAVY); tb(s, x, 2.0, wd, 0.6, h, 15, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
rows = [('1', '기내 와이파이', '+11.8%p', '모든 승객 유형에 효과\n(개인 여행 승객 +30%p, 출장 +7%p)', ORG),
        ('2', '온라인 탑승', '+8.2%p', '출장 승객에서 효과 집중 (+11%p)', ORG),
        ('후보', '다리 공간 · 체크인', '+3.4%p · +2.6%p', '1·2순위 이후 검토', GRY),
        ('제외', '음식 · 게이트 위치', '거의 없음', '투자 후순위', GRY)]
for i, (a, b, c, d, col) in enumerate(rows):
    y = 2.7 + i * 0.85
    for x, wd, v in zip([0.65, 1.95, 5.15, 8.15], [1.25, 3.1, 2.95, 4.5], [a, b, c, d]):
        rect(s, x, y, wd, 0.78, LIGHT)
    tb(s, 0.65, y, 1.25, 0.78, a, 18, True, col, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, 1.95, y, 3.1, 0.78, b, 18, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, 5.15, y, 2.95, 0.78, c, 18, True, col, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, 8.25, y, 4.3, 0.78, d.split('\n'), 13, False, INK, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
note(s, '→ 승객 구분 없이 와이파이·온라인 탑승부터 전체 개선', 6.35)
tb(s, 0.65, 7.05, 11.5, 0.35, '※ 상관 기반 추정이며 실제 효과는 시범 적용으로 확인 필요', 11, False, GRY)

# 13 한계
s = slide('한계와 향후 과제', 13, 3)
for i, (h, items, col) in enumerate([('한계', ['설문 상관 분석 → 인과는 시범 적용(A/B)으로 확인 필요', '0점(해당없음)을 낮은 점수와 동일 처리', '개선 비용 데이터 없음 → ROI 산출 불가'], GRY),
                                     ('향후 과제', ['Gradient Boosting + SHAP로 상호작용 확인', '개선 비용 시나리오 반영한 ROI 우선순위', '대시보드화 (Tableau/Streamlit)'], ORG)]):
    x = 0.65 + i * 6.15
    rect(s, x, 2.1, 5.85, 4.2, LIGHT); rect(s, x, 2.1, 5.85, 0.8, col)
    tb(s, x, 2.1, 5.85, 0.8, h, 22, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, x + 0.3, 3.2, 5.3, 3.0, ['• ' + t for t in items], 17, False, INK)


prs.save(str(OUTPUT / 'airline_satisfaction_deck.pptx'))
